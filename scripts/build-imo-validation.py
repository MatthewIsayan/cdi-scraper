import json, csv, re, hashlib
from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/imo-validation'; OUT.mkdir(parents=True,exist_ok=True)
DATA=json.loads((ROOT/'exports/aguilar-review/refined-data.json').read_text())
DATE='2026-10-08'
QUESTIONS=[
('Completeness','Did you review every surname-search result, or was the approved list a handpicked subset? Were any first-name initials, cities, or pages skipped?'),
('Life authority','Is active general Life authority mandatory? Should funeral/burial-limited producers have a separate recruiting pool?'),
('Company exclusions','Which exact State Farm/Farmers relationships cause rejection? Do you exclude affiliates such as Foremost, or only named relationships? Approved examples have Foremost appointments.'),
('Other channels','Are bank/securities, New York Life career, Northwestern, Knights of Columbus, Thrivent, Allstate, and benefits-broker profiles rejected, conditional, or acceptable? What evidence establishes the rule?'),
('Contact and territory','Is missing contact information a rejection or an enrichment task? Are central/toll-free numbers acceptable? Is California residency required? Ramiro (4192130) is an approved nonresident exception.'),
('Product mix','Do you exclude health-focused or P&C-focused producers, or accept anyone with active Life authority? What evidence must establish final expense, IUL, or fixed-annuity activity?'),
('Other filters','Do tenure, variable authority, renewal timing, agency/sponsor relationships, or licensing orders affect the decision? State each threshold and exception.'),
('Corrections','For every omitted active-Life producer: approve now, reject with the actual reason, defer with the missing fact, or mark not previously reviewed. Start with the 35 unexplained first-call candidates.')]

RULES=[
['L1','General Life authority','FAIL if verified record lacks active general Life','Qualification gate for stated target; reviewer must ratify','0 / 1347; 20 unverifiable records REVIEW','No general Life is different from missing/unavailable data; limited burial authority separate.'],
['C1','State Farm / Farmers','FAIL for established current named relationship','User-confirmed exclusion','0 / 17','Confirm exact company scope; do not expand automatically to corporate affiliates.'],
['R1','Bank / securities agency','REVIEW actual role and outside contracting','Inferred channel concern','0 / 20','A carrier appointment alone does not prove an exclusive role.'],
['R2','Career / member channel','REVIEW actual contract and portability','Inferred channel concern','0 / 10','New York Life, NYLIFE, Northwestern, Knights, Thrivent signals are not contracts.'],
['R3','Benefits broker / Allstate','REVIEW role, product business, distribution','Inferred channel concern','Benefits: 0 / 6','Named-marker heuristic is incomplete; no blanket Allstate ban confirmed.'],
['R4','Health / P&C emphasis','REVIEW product focus when Life business unclear','Inferred product concern','Health-only: 0 / 4; health-skew: 2 / 13; P&C-skew: 2 / 20','Approved exceptions prohibit universal rejection. Carrier mix does not prove sales.'],
['R5','No / indirect business contact','REVIEW or enrich contact','Operational hold; rejection unconfirmed','Current phone: 135 / 93','Approved PDF provides one missing-current-phone contact; no shared-number ban.'],
['R6','Outside California baseline','REVIEW geographic scope','Pattern with approved exception','CA resident: 135 / 153','Nonresident approval exists; state/address discrepancies must be resolved.'],
['R7','Licensing orders','REVIEW actual order and contracting effect','Observed order; effect unknown','Order-review signal: 1 / 2','Removal orders and resolved history are not current restrictions.'],
['S1','WFG / Primerica','Positive segment signal; not automatic PASS','Observed approved target groups','WFG: 9 / 9; Primerica: 42 / 20','Original PDF headings identify fewer than current relationships.'],
['S2','Mixed P&C / variable authority','Retain; assess surrounding profile','Blanket exclusion contradicted','Active P&C: 6 / 44; variable: 4 / 31','These authorities alone do not establish poor fit.'],
['S3','Tenure / expiry / languages / sponsors','Outreach context; no automatic FAIL','No threshold confirmed','Sponsor: 1 / 14','Tenure is issue-date tenure, not age or continuous production.'],
['P1','No unresolved screen issue','PASS initial screen; preserve original label','Provisional screening result','Computed in workbook','PASS is neither reviewer approval nor contracting eligibility.']]

def textrows(items): return '\n'.join(' / '.join(str(x) for x in a) for a in items)
def active_names(r):return '; '.join(q['qualification'] for q in r['qualifications'] if q['status']=='Active')
def issues(r):
    out=[]
    for key,rule,explain in [('bank_signal','R1','Bank/securities role and outside contracting'),('career_signal','R2','Career/member channel and portability'),('benefits_agency_signal','R3','Benefits-broker role and target-product activity'),('health_only_appointments','R4','A&H/disability appointments only'),('health_skew','R4','Selected health-carrier Life mix'),('pc_skew','R4','P&C emphasis without separate core-Life signal'),('order_review','R7','Actual order terms and carrier contracting')]:
        if r[key]:out.append((rule,explain))
    if any(a[0].startswith('ALLSTATE ') for a in r['appointments']):out.append(('R3','Allstate distribution role'))
    p=re.sub(r'\D','',r['effective_business_phone'])
    if not p:out.append(('R5','No collected public business phone'))
    elif p.startswith(('800','888','877','866','855','844','833')):out.append(('R5','Toll-free number: direct reachability'))
    elif r['phone_source']!='Current CDI':out.append(('R5','PDF-only contact: currency verification'))
    if not r['flags']['ca_resident']:out.append(('R6','Geographic scope'))
    return out

prepared=[]
for r in DATA:
    ri=issues(r) if r['active_life'] else []
    if not r['qualifications']:decision='REVIEW';basis='L1';why='Detail page displays no qualification information. Life authority cannot be verified; missing data is not a qualification failure. Historical omission reason unknown.'
    elif not r['active_life']:decision='FAIL';basis='L1';why='Verified qualifications do not include active general Life. Current target-pool mismatch; actual historical omission reason unknown.'
    elif r['flags']['state_farm_or_farmers']:decision='FAIL';basis='C1';why='Current State Farm/Farmers relationship matches the user-confirmed exclusion. Reviewer must verify that this was the historical reason.'
    elif ri:decision='REVIEW';basis=', '.join(dict.fromkeys(a[0] for a in ri));why='Requires verification: '+'; '.join(a[1] for a in ri)+'. These facts do not establish the original reviewer’s motive.'
    else:decision='PASS';basis='P1';why='Passes the provisional observable screen. Actual business activity and willingness are unknown; historical omission remains unexplained.' if not r['approved'] else 'Passes provisional observable screen; approval is independently established by PDF membership.'
    lifeq=next((q for q in r['qualifications'] if q['qualification']=='Life'),{})
    question=('Why was this person approved despite: '+'; '.join(a[1] for a in ri)+'? Confirm the exception or any intervening record change.' if r['approved'] and ri else 'Confirm approval criteria and any unrecorded reason.' if r['approved'] else 'Was this person reviewed? If rejected, give the actual rule and evidence. If acceptable, correct the omission. '+('Resolve: '+'; '.join(a[1] for a in ri)+'.' if ri else 'Confirm '+('the Life-qualification gate and any record change since selection.' if not r['active_life'] else 'the named company exclusion.' if decision=='FAIL' else 'whether there was an unrecorded filter or a missed search result.')))
    evidence_path=ROOT/'exports/aguilar-review/details'/f"{r['license_number']}.html"
    d={'name':r['current_name'],'license':r['license_number'],'original':'Approved' if r['approved'] else 'Omitted','decision':decision,'rule_ids':basis,'why':why,'review_question':question,'reviewer_decision':'Pending','actual_reason':'','reviewer_evidence':'','reviewer_name_date':'','life_status':lifeq.get('status','Not displayed'),'life_issued':lifeq.get('issued',''),'life_expires':lifeq.get('expires',''),'other_active':active_names(r),'phone':r['effective_business_phone'],'phone_source':r['phone_source'],'address':r['business_address'],'residency':r['residency'],'search_state':r['search_state'],'agency':textrows(r['agencies']),'appointments':textrows(r['appointments']),'orders':textrows(r['enforcement']),'complaints':textrows(r['complaints']),'sponsors':textrows(r['sponsor_relationships']),'employers':textrows(r['employer_relationships']),'languages':r['languages'],'tenure':r['life_tenure_years'],'expiry_days':r['life_expiration_days'],'recent_appointment':r['most_recent_life_insurer_appointment'],'positive':r['positive_signals'],'flags':r['review_flags'],'original_pdf_group':r.get('pdf_group',''),'pdf_phone':r.get('pdf_phone',''),'pdf_name':r.get('name',''),'aliases':' | '.join(r['search_names']),'source':r['url'],'check_date':DATE,'cache_sha256':hashlib.sha256(evidence_path.read_bytes()).hexdigest(),'prior_tier':r['priority_tier'],'prior_route':r['proposed_route'],'historical_reason_status':'Not supplied; membership is known, motive is not','active_life':r['active_life']}
    d['all_qualifications']='\n'.join(' / '.join(str(q.get(k,'')) for k in ['qualification','status','issued','status_date','expires'])for q in r['qualifications'])
    d['search_status']=r['search_status']
    d['name_source']='Current detail page' if r['current_name'] else 'Search result only; detail-page name blank'
    if not d['name']:d['name']=' '.join(r['search_names'][0].split())
    if not r['qualifications']:d['review_question']='Was this record reviewed? The detail page provides no qualifications. Verify the actual license type and Life authority before deciding; do not reject for missing data.'
    prepared.append(d)
prepared.sort(key=lambda d:(not d['active_life'],d['original']=='Approved',d['prior_tier'],d['name'],d['license']))
(OUT/'validation-data.json').write_text(json.dumps({'records':prepared,'rules':RULES,'questions':QUESTIONS},ensure_ascii=False,indent=2),encoding='utf-8')
counts=Counter(d['decision'] for d in prepared if d['active_life']);counts_orig=Counter((d['original'],d['decision']) for d in prepared if d['active_life'])
print('Strict playbook decisions:',dict(counts),'by membership:',dict(counts_orig))

def make_doc(title):
    d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(6)
    d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.line_spacing=1.12
    d.styles['Title'].font.size=Pt(22);d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(12)
    footer=s.footer.paragraphs[0];footer.text='IMO recruiting validation • Evidence checked October 8, 2026 • '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for run in footer.runs:run.font.size=Pt(8)
    d.add_heading(title,0);return d
def p(d,t):d.add_paragraph(t)
def h(d,t):d.add_heading(t,1)
def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for i,v in enumerate(row):t.add_row() if False else None
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=1;pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF3');pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
                for run in para.runs:run.font.size=Pt(9);run.bold=ri==0
    d.add_paragraph('');return t
def page(d,title):d.add_page_break();h(d,title)

d=make_doc('IMO recruiting screening experiment')
p(d,'Full findings, unresolved omissions, and reviewer validation protocol\nLife • Final expense • IUL • Fixed annuities')
h(d,'What the evidence establishes')
p(d,'The supplied PDFs contain 136 approved licenses. The reconstructed California Department of Insurance (CDI) search contains 1,705 unique licenses. All 136 approved licenses were found. Of the 1,569 omitted licenses, 202 currently show active general Life authority and need the most careful review; 1,367 do not show that authority.')
table(d,['Cohort','Count','Interpretation'],[['Approved in supplied PDFs',136,'Known membership; selection motive undocumented'],['Omitted, active general Life',202,'Full person-by-person reviewer work queue'],['Omitted, other records',1367,'1,347 qualification mismatches; 20 unverifiable qualification pages'],['Total reconstructed search',1705,'A–Z first-initial search; deduplicated by license']], [2.5,.6,3.8])
p(d,'Thirty-five omitted active-Life producers meet the earlier first-call profile. That is a useful discrepancy set, not proof of 35 mistakes. The original list could be incomplete, the reviewer could have applied an undocumented filter, or the public records could have changed.')
h(d,'Confidence and practical limits')
p(d,'Confidence is high in the recorded list membership and the extracted, dated public fields. Confidence is limited in explanations of the reviewer’s intent. No percentage of historical decision accuracy can be justified yet. State Farm/Farmers is the only company exclusion explicitly confirmed by the user; the active-Life gate follows the stated recruiting target and needs reviewer ratification.')
p(d,'The workbook contains every license, direct source links, qualification dates/statuses, contacts, agencies, line-specific appointments, orders, sponsors, all observed concerns, a provisional screening result, and editable reviewer answers. An omitted label is preserved even when the provisional result is PASS.')

page(d,'How the test cohort was reconstructed')
p(d,'The source is CDI’s public license lookup: https://cdicloud.insurance.ca.gov/cal/ . Individual name search: https://cdicloud.insurance.ca.gov/cal/IndividualNameSearch . Each workbook row links to its exact LicenseDetail page; license number and individual ID must both be preserved.')
p(d,'A surname-only search exceeded the result limit. Searches using first-name initials A through Z returned 2,035 rows. Historical names and multiple matching aliases were deduplicated to 1,705 licenses. Cached detail pages were retrieved and checked for identity on October 8, 2026. No retrieval failures remain in this dataset. This establishes coverage of the reconstructed A–Z search, not a guarantee against unusual first-name initials or every possible search edge case.')
p(d,'The original PDFs are dated October 5, 2026: Aguilar_Producers_Full_Detailed_Directory.pdf and Aguilar_Life_Producers_Recruiting_List.pdf. They are approval references, not instructions or proof of how the original search was performed. Their headings label 7 WFG, 33 Primerica, and 96 affiliation-unidentified producers. Current relationships identify WFG for 9 approved and Primerica for 42 approved. Unidentified does not mean independent.')
h(d,'Three distinct quantities')
p(d,'Historical label: approved or omitted from the supplied PDFs. Public fact: the qualification/contact/relationship displayed on the check date. Screening judgment: PASS, FAIL, or REVIEW under the provisional playbook. Keep all three separately; none substitutes for another.')
p(d,'Current CDI facts may differ from facts on the selection date. One approved producer has a PDF phone but no current displayed phone. One approved producer is currently nonresident. Twenty detail pages provide no qualification information and require verification; one bail-agent detail page has a blank name, so the workbook uses its search-result name and labels that source. A retrieved page is not necessarily complete evidence.')
h(d,'What license pages cannot establish')
p(d,'They do not establish sales volume, actual final-expense/IUL/fixed-annuity production, interest in joining, contract exclusivity, release requirements, ownership of a book, or carrier appointment eligibility for a new IMO. A carrier name or Life appointment is an alignment signal, not product-level proof. Publicly reported languages are retained; no language or ethnicity is inferred from a surname.')

page(d,'Observed factors and what they do—and do not—explain')
p(d,'Counts below compare all 136 approved active-Life producers with all 202 omitted active-Life producers. Factors overlap and must not be summed as distinct rejection reasons.')
factors=[('Current CDI phone',135,93,'Operational reachability; one approved PDF-only phone'),('California resident baseline',135,153,'Strong pattern, but approved nonresident exception'),('State Farm / Farmers',0,17,'User-confirmed exclusion'),('WFG relationship',9,9,'Positive target segment; omissions remain'),('Primerica relationship',42,20,'Positive target segment; omissions remain'),('Bank / securities agency',0,20,'Channel hypothesis; no proven exclusive contract'),('Career / member-channel marker',0,10,'Contract/portability verification'),('Benefits-broker agency marker',0,6,'Retail target-product activity unestablished'),('Only A&H/disability appointments',0,4,'Life authority exists; displayed appointments differ'),('Selected health-carrier-only Life mix',2,13,'Approved exceptions; product inquiry'),('P&C-skew heuristic',2,20,'Approved exceptions; ask actual Life activity'),('Active Property/Casualty authority',6,44,'Blanket exclusion contradicted'),('Variable authority',4,31,'Blanket exclusion contradicted'),('Order requiring further review',1,2,'Actual current terms need inspection'),('Sponsor relationship',1,14,'Qualification-specific; not a universal employer'),('No agency or insurer displayed',25,32,'Business activity and independence unknown')]
table(d,['Factor','Approved','Omitted','Meaning'],factors,[2.2,.85,.85,3])
p(d,'Marker sets are transparent heuristics, not an exhaustive industry taxonomy. “Health-only” means appointments exclusively for A&H/disability; it must not be conflated with simply having no direct Life insurer appointment. Tenure in the reachable CA-Life cohort has medians of 2.5 years approved and 7.0 years omitted; this does not establish a cutoff or causation.')

page(d,'Provisional pass / fail playbook')
for title,body in [
('1. Verify identity and general Life authority','Use license number; inspect the Life qualification’s own status. Verified absence of active general Life is FAIL for this pool. Unavailable or unverifiable records are REVIEW. Funeral/burial-limited authority gets a separate pool if requested.'),
('2. Apply the confirmed company exclusion','Established current State Farm/Farmers relationships are FAIL. Confirm company scope with the reviewer. Foremost alone is not an automatic Farmers rejection because approved examples have Foremost appointments.'),
('3. Resolve contact, geography, product, channel, and order issues','Use REVIEW for missing/indirect contact, outside-California scope, bank/career/benefits/Allstate channels, unclear health/P&C product focus, and material orders. Record every applicable issue; do not present any of them as a confirmed historical reason.'),
('4. Pass the observable screen','Active general Life, usable contact, suitable geography, no confirmed exclusion, and no unresolved screen issue produces provisional PASS. Keep actual PDF approval separately. No displayed affiliation is not itself failure; ask about business activity.'),
('5. Keep context out of hard exclusions','P&C or variable authority, tenure, upcoming renewal, languages, carrier count, and qualification-specific sponsors do not alone cause FAIL. Resolved disciplinary history is not automatically a current restriction.')]:
    d.add_heading(title,2);p(d,body)
table(d,['Strict playbook result','Approved','Omitted active Life'],[[k,counts_orig[('Approved',k)],counts_orig[('Omitted',k)]]for k in ['PASS','REVIEW','FAIL']],[3.5,1.5,1.9])
p(d,'The strict playbook is more conservative than the earlier first-call ranking: it holds PDF-only contacts and unresolved distribution details for REVIEW. The 35 first-call candidates remain the discrepancy queue, not a certified or newly approved list. The workbook includes both results so this difference is visible.')

page(d,'Approved controls that prevent over-rejection')
p(d,'Audit approved people with apparent concerns before turning a correlation into a hard rule. These eight approvals fall outside the earlier normal first-call profile; their actual approvals must remain intact.')
exceptions=[r for r in DATA if r['approved'] and r['priority_tier']!=1]
table(d,['Approved person / license','Observed issue requiring an explanation'],[[r['current_name']+'\n'+r['license_number'],r['main_reason']]for r in exceptions],[2.6,4.3])
p(d,'Additional controls: six approved producers have active Property/Casualty authority; four have variable authority. Jessica Alejandra Aguilar (0J10393) has P&C relationships plus American General and Southwest Life. Approved AAA, Guardian, health-carrier, and removed-restriction examples preclude treating those observations as universal disqualifiers.')
p(d,'For each control, ask whether the trait is acceptable, conditionally acceptable, an explicit exception, or a later record change. If the proposed rule rejects a positive control, revise the rule or document the reviewer-authorized exception—do not silently relabel the approved person.')

first=[r for r in DATA if not r['approved'] and r['priority_tier']==1]
first.sort(key=lambda r:r['current_name'])
for part,start in enumerate(range(0,len(first),12),1):
    page(d,f'Unexplained first-call omissions • {part} of 3')
    p(d,'Each person below has active general Life, a collected business phone, a California baseline, and no impediment in the earlier first-call ranking. The reviewer must supply the actual omission reason or correct membership. These are hypotheses to investigate, not new approvals.')
    rows=[]
    for r in first[start:start+12]:
        sig='WFG' if r['flags']['wfg'] else 'Primerica' if r['flags']['primerica'] else 'No displayed affiliation' if r['flags']['no_agency_or_insurer_displayed'] else 'Direct Life portfolio; inspect agency and product mix'
        caveat='Mixed P&C; approved analogues exist' if r['flags']['active_property_or_casualty'] else 'Variable authority; not an exclusion' if r['variable_authority'] else 'Actual target-product production unknown'
        rows.append([r['current_name']+'\n'+r['license_number'],sig+'\n'+caveat])
    table(d,['Omitted person / license','Useful signal and unresolved point'],rows,[3,3.9])
    p(d,'The workbook provides exact appointment names, dates, contact information, source links, all additional flags, and a reviewer decision field for these and every other producer.')

page(d,'How to resolve all omissions and validate the rules')
for title,body in [
('Adjudicate membership before modeling it','First establish whether all search results were actually reviewed. “Not previously reviewed” is a distinct answer, not a rejection. If the original search universe was smaller, record its boundaries; omitted positives can then be search-coverage misses instead of policy exceptions.'),
('Review the entire active-Life queue','Fill the amber fields in “Life review” for all 338 rows: 202 omissions plus 136 approved controls. Start with the 35 unexplained first-call candidates, then other omissions, then approved exceptions. Reviewer choices: Approve, Reject, Defer, or Not reviewed. Every Reject needs the actual reason and evidence; every Defer needs the missing fact.'),
('Confirm the other-license group without inventing motives','The “Other licenses” sheet includes all 1,367 records: 1,347 verified qualification mismatches plus 20 records with no qualification information. Resolve those 20 REVIEW records individually. The reviewer can affirm the mismatch gate in bulk after checking representative and limited-authority records. Preserve unknown historical motive and explain individual overrides.'),
('Turn answers into testable rules','For each confirmed reason, record rule ID, triggering fields, decision, scope, threshold if any, exceptions, confirmation date, and examples. Separate eligibility failure from contact enrichment, territory hold, and product/channel inquiry. Never use an inferred exclusive contract as a hard rule.'),
('Reconcile contradictions','An omitted producer acceptable under confirmed rules becomes a correction candidate. An approved producer failing a proposed rule requires a rule revision, explicit exception, or dated change explanation. No unexplained rejected row should be counted as successfully modeled.'),
('Test on a new cohort','Freeze this dataset and the adjudicated rule version. Use a different surname as a prospective holdout and have the reviewer label it independently before seeing the model result. Report agreement, false rejections, false approvals, review rate, and unresolved cases using only adjudicated labels. Do not report the current in-sample ranking as predictive accuracy.')]:
    d.add_heading(title,2);p(d,body)
p(d,'Ready-to-scale condition: all active-Life omissions and positive-control conflicts resolved; rule definitions approved; a new-cohort test completed; and every future exclusion traceable to dated evidence and a confirmed rule. Unresolved cases remain REVIEW.')

page(d,'Short question list for the original reviewer')
p(d,'Please use these questions to confirm the playbook, then return the completed workbook. Our objective is to reproduce your intended screen and identify any missed acceptable producers, rather than rationalize the list after the fact.')
for i,(title,q) in enumerate(QUESTIONS,1):d.add_heading(f'{i}. {title}',2);p(d,q)
p(d,'For each person, return: Approve / Reject / Defer / Not reviewed; the actual reason; the source or unrecorded fact; and your name/date. Please explicitly identify corrected omissions and authorized exceptions. No outreach or messages have been sent as part of this analysis.')
d.save(OUT/'IMO_Recruiting_Analysis.docx')

c=make_doc('Reviewer checklist')
p(c,'Aguilar test cohort • Evidence checked October 8, 2026\nTarget: Life, final expense, IUL, and fixed annuities')
p(c,'136 approved • 202 omitted with active general Life • 1,367 other omitted licenses. Start with the 35 unexplained first-call candidates; also verify the 20 other records with missing qualification information.')
for i,(title,q) in enumerate(QUESTIONS,1):
    para=c.add_paragraph();rr=para.add_run(f'{i}. {title}: ');rr.bold=True;para.add_run(q)
    para.paragraph_format.space_after=Pt(8)
p(c,'Per-person response: Approve / Reject / Defer / Not reviewed + actual reason + evidence + reviewer/date. Correct any acceptable omissions. Explain approved exceptions before adding a hard exclusion.')
p(c,'Current safe screen: verify active general Life; exclude confirmed State Farm/Farmers relationships; REVIEW unresolved channel/product/contact/territory/orders; PASS only when no screen issue remains. PASS does not mean historical approval.')
c.save(OUT/'Reviewer_Checklist.docx')
print('Created report and checklist; prepared',len(prepared),'review records.')
